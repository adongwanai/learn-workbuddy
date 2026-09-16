#!/usr/bin/env python3
"""Small, offline XLSX helper for the learn-workbuddy Excel Skill."""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


def fail(message: str) -> None:
    raise ValueError(message)


def emit(payload: dict[str, Any]) -> None:
    print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))


def require_xlsx(path: Path) -> None:
    if path.suffix.lower() != ".xlsx":
        fail(f"仅支持 .xlsx 文件: {path}")


def load_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        fail(f"JSON 文件不存在: {path}")
    except json.JSONDecodeError as exc:
        fail(f"JSON 格式错误: {path}: {exc}")
    if not isinstance(value, dict):
        fail(f"JSON 顶层必须是对象: {path}")
    return value


def parse_date(value: Any) -> date:
    if not isinstance(value, str):
        fail("date 必须是 YYYY-MM-DD 字符串")
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        fail(f"日期格式错误: {value}: {exc}")


def typed_value(value: Any) -> Any:
    if not isinstance(value, dict):
        return value
    keys = set(value)
    if keys == {"date"}:
        return parse_date(value["date"])
    if keys == {"formula"}:
        formula = value["formula"]
        if not isinstance(formula, str) or not formula.startswith("="):
            fail("formula 必须是以 '=' 开头的字符串")
        return formula
    if keys == {"value"}:
        return value["value"]
    fail("单元格对象只能包含 date、formula 或 value 其中一个字段")


def workbook_summary(path: Path) -> dict[str, Any]:
    require_xlsx(path)
    if not path.exists():
        fail(f"工作簿不存在: {path}")
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        fail(f"无法打开工作簿: {path}: {exc}")

    sheets = []
    formula_count = 0
    date_count = 0
    for sheet in workbook.worksheets:
        sheet_formulas = 0
        sheet_dates = 0
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    sheet_formulas += 1
                if isinstance(cell.value, (date, datetime)):
                    sheet_dates += 1
        formula_count += sheet_formulas
        date_count += sheet_dates
        sheets.append(
            {
                "title": sheet.title,
                "state": sheet.sheet_state,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                "formula_count": sheet_formulas,
                "date_count": sheet_dates,
            }
        )

    return {
        "path": str(path),
        "format": "xlsx",
        "sheet_count": len(workbook.worksheets),
        "sheets": sheets,
        "formula_count": formula_count,
        "date_count": date_count,
        "has_formulas": formula_count > 0,
    }


def write_cell(cell, value: Any) -> None:
    cell.value = typed_value(value)


def create_workbook(output: Path, spec: dict[str, Any]) -> dict[str, Any]:
    require_xlsx(output)
    if output.exists():
        fail(f"输出文件已存在，不覆盖已有文件: {output}")
    sheet_name = spec.get("sheet")
    headers = spec.get("headers")
    rows = spec.get("rows", [])
    if not isinstance(sheet_name, str) or not sheet_name.strip():
        fail("spec.sheet 必须是非空字符串")
    if not isinstance(headers, list) or not all(isinstance(item, str) for item in headers):
        fail("spec.headers 必须是字符串数组")
    if not isinstance(rows, list) or not all(isinstance(row, list) for row in rows):
        fail("spec.rows 必须是二维数组")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)
    for row_number, row in enumerate(rows, start=2):
        if len(row) != len(headers):
            fail(f"第 {row_number} 行列数与 headers 不一致")
        for column, value in enumerate(row, start=1):
            write_cell(sheet.cell(row=row_number, column=column), value)

    widths = spec.get("column_widths", {})
    if widths:
        if not isinstance(widths, dict):
            fail("column_widths 必须是对象")
        for column, width in widths.items():
            if not isinstance(column, str) or not isinstance(width, (int, float)):
                fail("column_widths 必须是列名到数字宽度的映射")
            sheet.column_dimensions[column].width = width

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return workbook_summary(output)


def edit_workbook(input_path: Path, output_path: Path, operations: dict[str, Any]) -> dict[str, Any]:
    require_xlsx(input_path)
    require_xlsx(output_path)
    if not input_path.exists():
        fail(f"输入工作簿不存在: {input_path}")
    if input_path.resolve() == output_path.resolve():
        fail("编辑时输入文件和输出文件必须是不同路径")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(input_path, output_path)
    try:
        workbook = load_workbook(output_path, data_only=False, read_only=False)
    except Exception as exc:
        output_path.unlink(missing_ok=True)
        fail(f"无法打开输入工作簿: {exc}")

    for item in operations.get("set_cells", []):
        if not isinstance(item, dict):
            fail("set_cells 中的每一项必须是对象")
        sheet_name = item.get("sheet", "")
        sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else None
        if sheet is None:
            fail(f"目标 Sheet 不存在: {item.get('sheet')}")
        cell_name = item.get("cell")
        if not isinstance(cell_name, str):
            fail("set_cells.cell 必须是字符串")
        value_fields = [key for key in ("value", "date", "formula") if key in item]
        if len(value_fields) != 1:
            fail("set_cells 必须且只能提供 value、date 或 formula 之一")
        key = value_fields[0]
        value = {"date": item[key]} if key == "date" else {"formula": item[key]} if key == "formula" else item[key]
        write_cell(sheet[cell_name], value)

    for item in operations.get("append_rows", []):
        if not isinstance(item, dict) or not isinstance(item.get("values"), list):
            fail("append_rows 每项必须包含 values 数组")
        sheet_name = item.get("sheet")
        if sheet_name not in workbook.sheetnames:
            fail(f"目标 Sheet 不存在: {sheet_name}")
        sheet = workbook[sheet_name]
        row_number = sheet.max_row + 1
        for column, value in enumerate(item["values"], start=1):
            write_cell(sheet.cell(row=row_number, column=column), value)

    for item in operations.get("number_formats", []):
        if not isinstance(item, dict) or not isinstance(item.get("format"), str):
            fail("number_formats 每项必须包含 format 字符串")
        sheet_name = item.get("sheet")
        if sheet_name not in workbook.sheetnames:
            fail(f"目标 Sheet 不存在: {sheet_name}")
        workbook[sheet_name][item["cell"]].number_format = item["format"]

    workbook.save(output_path)
    return workbook_summary(output_path)


def cell_value_matches(actual: Any, expected: Any) -> bool:
    if isinstance(expected, dict) and set(expected) == {"date"}:
        return isinstance(actual, (date, datetime)) and actual.date() == parse_date(expected["date"]) if isinstance(actual, datetime) else actual == parse_date(expected["date"])
    if isinstance(expected, dict) and set(expected) == {"formula"}:
        return actual == expected["formula"]
    return actual == expected


def validate_workbook(path: Path, expectations: dict[str, Any]) -> dict[str, Any]:
    summary = workbook_summary(path)
    workbook = load_workbook(path, data_only=False, read_only=False)
    failures = []

    expected_sheets = expectations.get("sheets", [])
    for sheet_name in expected_sheets:
        if sheet_name not in workbook.sheetnames:
            failures.append(f"缺少 Sheet: {sheet_name}")

    for reference, expected in expectations.get("cells", {}).items():
        if "!" not in reference:
            failures.append(f"单元格引用缺少 '!': {reference}")
            continue
        sheet_name, cell_name = reference.split("!", 1)
        if sheet_name not in workbook.sheetnames:
            failures.append(f"单元格引用的 Sheet 不存在: {reference}")
            continue
        actual = workbook[sheet_name][cell_name].value
        if not cell_value_matches(actual, expected):
            failures.append(f"{reference}: 期望 {expected!r}，实际 {actual!r}")

    for reference in expectations.get("required_formulas", []):
        if "!" not in reference:
            failures.append(f"公式引用缺少 '!': {reference}")
            continue
        sheet_name, cell_name = reference.split("!", 1)
        if sheet_name not in workbook.sheetnames:
            failures.append(f"公式引用的 Sheet 不存在: {reference}")
            continue
        value = workbook[sheet_name][cell_name].value
        if not isinstance(value, str) or not value.startswith("="):
            failures.append(f"{reference}: 不是公式")

    result = {
        "status": "passed" if not failures else "failed",
        "path": str(path),
        "summary": summary,
        "failures": failures,
    }
    if failures:
        raise ValueError(json.dumps(result, ensure_ascii=False))
    return result


def recalculate_workbook(path: Path) -> dict[str, Any]:
    require_xlsx(path)
    if not path.exists():
        fail(f"工作簿不存在: {path}")
    workbook = load_workbook(path, data_only=False, read_only=False)
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(path)

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return {
            "status": "unavailable",
            "path": str(path),
            "reason": "未检测到 LibreOffice；已设置 Excel/WPS 打开时自动重算标记",
        }

    with tempfile.TemporaryDirectory(prefix="workbuddy-xlsx-") as temp_dir:
        temp_path = Path(temp_dir)
        command = [
            soffice,
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(temp_path),
            str(path),
        ]
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
        recalculated = temp_path / path.name
        if completed.returncode != 0 or not recalculated.exists():
            return {
                "status": "failed",
                "path": str(path),
                "reason": (completed.stdout + completed.stderr).strip(),
            }
        shutil.copy2(recalculated, path)
    return {"status": "recalculated", "path": str(path)}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="WorkBuddy 基础 XLSX 工作簿工具")
    subparsers = parser.add_subparsers(dest="command", required=True)

    inspect_parser = subparsers.add_parser("inspect", help="检查工作簿结构")
    inspect_parser.add_argument("input", type=Path)

    create_parser = subparsers.add_parser("create", help="创建工作簿")
    create_parser.add_argument("output", type=Path)
    create_parser.add_argument("--spec", required=True, type=Path)

    edit_parser = subparsers.add_parser("edit", help="复制并编辑工作簿")
    edit_parser.add_argument("input", type=Path)
    edit_parser.add_argument("output", type=Path)
    edit_parser.add_argument("--operations", required=True, type=Path)

    validate_parser = subparsers.add_parser("validate", help="验证工作簿")
    validate_parser.add_argument("input", type=Path)
    validate_parser.add_argument("--expect", required=True, type=Path)

    recalc_parser = subparsers.add_parser("recalc", help="设置或执行公式重算")
    recalc_parser.add_argument("input", type=Path)
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        if args.command == "inspect":
            emit(workbook_summary(args.input))
        elif args.command == "create":
            emit(create_workbook(args.output, load_json(args.spec)))
        elif args.command == "edit":
            emit(edit_workbook(args.input, args.output, load_json(args.operations)))
        elif args.command == "validate":
            emit(validate_workbook(args.input, load_json(args.expect)))
        elif args.command == "recalc":
            emit(recalculate_workbook(args.input))
    except (OSError, ValueError) as exc:
        print(json.dumps({"status": "error", "error": str(exc)}, ensure_ascii=False))
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
