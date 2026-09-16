from __future__ import annotations

import importlib.util
import json
import os
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SKILL_DIR = ROOT / "skills" / "excel-workbook"
SKILL_MD = SKILL_DIR / "SKILL.md"
SCRIPT = SKILL_DIR / "scripts" / "workbook_tool.py"


def run_tool(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )


def json_output(result: subprocess.CompletedProcess[str]) -> dict:
    assert result.stdout, result.stderr
    return json.loads(result.stdout)


def load_s16():
    stub_dir = ROOT / "tests" / "stubs"
    sys.path.insert(0, str(stub_dir))
    saved_anthropic = sys.modules.pop("anthropic", None)
    old_model = os.environ.get("MODEL_ID")
    os.environ["MODEL_ID"] = "offline-test-model"
    module_name = "s16_excel_skill_loader_test"
    try:
        spec = importlib.util.spec_from_file_location(
            module_name, ROOT / "s16_skills_system" / "code.py"
        )
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        sys.modules[module_name] = module
        spec.loader.exec_module(module)
        return module
    finally:
        sys.path.remove(str(stub_dir))
        sys.modules.pop("anthropic", None)
        if saved_anthropic is not None:
            sys.modules["anthropic"] = saved_anthropic
        if old_model is None:
            os.environ.pop("MODEL_ID", None)
        else:
            os.environ["MODEL_ID"] = old_model


def test_skill_file_has_chinese_frontmatter_and_loads() -> None:
    text = SKILL_MD.read_text(encoding="utf-8")
    assert text.startswith("---\n")
    module = load_s16()
    frontmatter, body = module.parse_frontmatter(text)
    assert frontmatter["name"] == "excel-workbook"
    assert "Excel" in frontmatter["description"]
    assert frontmatter["allowed-tools"] == ["Bash"]
    assert "先检查，再修改" in body
    assert "workbook_tool.py" in body


def test_script_compiles_and_exposes_commands() -> None:
    compile_result = subprocess.run(
        [sys.executable, "-m", "py_compile", str(SCRIPT)],
        cwd=ROOT,
        capture_output=True,
        text=True,
    )
    assert compile_result.returncode == 0, compile_result.stderr
    help_result = run_tool("--help")
    assert help_result.returncode == 0, help_result.stderr
    assert all(
        command in help_result.stdout
        for command in ("inspect", "create", "edit", "validate", "recalc")
    )


def test_create_inspect_edit_validate_and_recalc(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    edited = tmp_path / "edited.xlsx"
    spec_path = tmp_path / "spec.json"
    operations_path = tmp_path / "operations.json"
    expectations_path = tmp_path / "expectations.json"

    spec_path.write_text(
        json.dumps(
            {
                "sheet": "销售数据",
                "headers": ["日期", "商品", "金额", "含税金额"],
                "rows": [
                    [
                        {"date": "2026-09-16"},
                        "A",
                        100,
                        {"formula": "=C2*1.1"},
                    ]
                ],
                "column_widths": {"A": 14, "C": 12},
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    created = run_tool("create", str(source), "--spec", str(spec_path))
    assert created.returncode == 0, created.stderr
    assert json_output(created)["sheet_count"] == 1

    inspected = run_tool("inspect", str(source))
    assert inspected.returncode == 0, inspected.stderr
    summary = json_output(inspected)
    assert summary["sheets"][0]["title"] == "销售数据"
    assert summary["formula_count"] == 1
    assert summary["date_count"] == 1

    operations_path.write_text(
        json.dumps(
            {
                "set_cells": [
                    {"sheet": "销售数据", "cell": "C2", "value": 120},
                    {"sheet": "销售数据", "cell": "D2", "formula": "=C2*1.1"},
                    {"sheet": "销售数据", "cell": "A2", "date": "2026-09-17"},
                ],
                "append_rows": [
                    {
                        "sheet": "销售数据",
                        "values": [
                            {"date": "2026-09-18"},
                            "B",
                            80,
                            {"formula": "=C3*1.1"},
                        ],
                    }
                ],
                "number_formats": [
                    {
                        "sheet": "销售数据",
                        "cell": "C2",
                        "format": "#,##0.00",
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    edited_result = run_tool(
        "edit",
        str(source),
        str(edited),
        "--operations",
        str(operations_path),
    )
    assert edited_result.returncode == 0, edited_result.stderr
    assert source.exists()
    assert edited.exists()

    workbook = load_workbook(edited, data_only=False)
    sheet = workbook["销售数据"]
    assert isinstance(sheet["A2"].value, (date, datetime))
    assert sheet["A2"].value.date() == date(2026, 9, 17)
    assert sheet["D2"].value == "=C2*1.1"
    assert sheet["D3"].value == "=C3*1.1"
    assert sheet["C2"].number_format == "#,##0.00"
    assert sheet["A1"].font.bold is True

    expectations_path.write_text(
        json.dumps(
            {
                "sheets": ["销售数据"],
                "cells": {
                    "销售数据!C2": 120,
                    "销售数据!D2": "=C2*1.1",
                    "销售数据!A2": {"date": "2026-09-17"},
                },
                "required_formulas": ["销售数据!D2", "销售数据!D3"],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    validated = run_tool(
        "validate",
        str(edited),
        "--expect",
        str(expectations_path),
    )
    assert validated.returncode == 0, validated.stderr
    assert json_output(validated)["status"] == "passed"

    recalculated = run_tool("recalc", str(edited))
    assert recalculated.returncode == 0, recalculated.stderr
    assert json_output(recalculated)["status"] in {"unavailable", "recalculated"}
    workbook = load_workbook(edited, data_only=False)
    assert workbook.calculation.calcMode == "auto"
    assert workbook.calculation.fullCalcOnLoad is True
    assert workbook.calculation.forceFullCalc is True


def test_failures_are_explicit(tmp_path: Path) -> None:
    missing = tmp_path / "missing.xlsx"
    inspected = run_tool("inspect", str(missing))
    assert inspected.returncode == 2
    assert "不存在" in json_output(inspected)["error"]

    source = tmp_path / "source.xlsx"
    source.write_bytes(b"not an xlsx")
    operations = tmp_path / "operations.json"
    operations.write_text("{}", encoding="utf-8")
    edited = run_tool(
        "edit",
        str(source),
        str(source),
        "--operations",
        str(operations),
    )
    assert edited.returncode == 2
    assert "不同路径" in json_output(edited)["error"]

    invalid_extension = run_tool("inspect", str(tmp_path / "file.xls"))
    assert invalid_extension.returncode == 2
    assert ".xlsx" in json_output(invalid_extension)["error"]


def test_validate_reports_mismatch(tmp_path: Path) -> None:
    workbook_path = tmp_path / "sample.xlsx"
    workbook = __import__("openpyxl").Workbook()
    workbook.active["A1"] = 1
    workbook.save(workbook_path)
    expectations = tmp_path / "expectations.json"
    expectations.write_text(
        json.dumps({"cells": {"Sheet!A1": 2}}, ensure_ascii=False),
        encoding="utf-8",
    )
    result = run_tool(
        "validate",
        str(workbook_path),
        "--expect",
        str(expectations),
    )
    assert result.returncode == 2
    assert json_output(result)["status"] == "error"
    assert "期望" in json_output(result)["error"]
